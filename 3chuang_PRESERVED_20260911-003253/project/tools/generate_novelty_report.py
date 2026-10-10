from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\cixing\3chuang")
OUT_DIR = ROOT / "output" / "word"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "甬柜易算项目科技查新报告（正式稿）.docx"
LOGO_PATH = ROOT / "tmp" / "pdfs" / "craic-asset-000.jpg"

PROJECT_NAME = "甬柜易算：基于数字孪生的宁波港出海企业装箱与合规 SaaS 引擎"
REPORT_DATE = "2026年7月11日"

ACCENT = "D9E2F3"
ACCENT_DARK = "1F4E78"
LIGHT_GRAY = "F2F2F2"
BORDER = "666666"


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is None:
        no_wrap = OxmlElement("w:noWrap")
        tc_pr.append(no_wrap)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_run_font(run, size=10.5, bold=False, color=None, east_asia="宋体", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_font(paragraph, size=10.5, bold=False, color=None, east_asia="宋体"):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold or bool(run.bold), color=color, east_asia=east_asia)


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    r_pr.append(sz)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9)
    return run


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def set_paragraph_box(paragraph, fill=ACCENT):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        p_bdr.append(el)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "100")
    spacing.set(qn("w:after"), "100")


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.widow_control = True

    for name, size, east_asia, color in (
        ("Heading 1", 15, "黑体", "000000"),
        ("Heading 2", 12.5, "黑体", "1F1F1F"),
        ("Heading 3", 11, "黑体", "1F1F1F"),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(20)
    title.font.bold = True
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def configure_section(section, cover=False):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.15 if not cover else 2.2)
    section.bottom_margin = Cm(1.9 if not cover else 2.2)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.75)
    section.gutter = Cm(0)


def clear_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for p in section.header.paragraphs:
        p.clear()
    for p in section.footer.paragraphs:
        p.clear()


def add_running_header(section):
    header = section.header
    header.is_linked_to_previous = False
    p0 = header.paragraphs[0]
    p0.clear()
    table = header.add_table(rows=1, cols=3, width=Cm(16.7))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.25, 12.15, 3.30]
    for i, w in enumerate(widths):
        set_cell_width(table.cell(0, i), w)
        set_cell_margins(table.cell(0, i), top=0, bottom=40, start=20, end=20)
        set_cell_border(table.cell(0, i), bottom={"val": "single", "sz": "8", "color": "7F7F7F", "space": "0"})
        table.cell(0, i).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if LOGO_PATH.exists():
        p = table.cell(0, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Cm(0.78), height=Cm(0.86))
    center = table.cell(0, 1).paragraphs[0]
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = center.add_run("第二十八届中国机器人及人工智能大赛")
    set_run_font(r, size=11, bold=True, east_asia="黑体")
    right = table.cell(0, 2).paragraphs[0]
    set_cell_no_wrap(table.cell(0, 2))
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right.add_run("人工智能创新赛")
    set_run_font(r, size=8.5, east_asia="宋体", color="404040")
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.line_spacing = 1


def add_running_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("- ")
    set_run_font(r, size=9)
    add_field(p, "PAGE", "1")
    r = p.add_run(" -")
    set_run_font(r, size=9)
    p.paragraph_format.first_line_indent = Pt(0)


def add_body_paragraph(container, text="", *, bold_lead=None, first_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10.5, space_after=0, keep=False):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Pt(21) if first_indent else Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True, east_asia="黑体")
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size)
    return p


def add_bullet(container, text, level=0):
    p = container.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(0.65 + level * 0.6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_section_heading(doc, text, *, page_break_before=False):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(text)
    set_run_font(r, size=15, bold=True, east_asia="黑体")
    set_paragraph_box(p)
    return p


def add_subheading(doc, text, level=2):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=12.5 if level == 2 else 11, bold=True, east_asia="黑体")
    return p


def add_table(doc, headers, rows, widths, font_size=9, header_fill=LIGHT_GRAY, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = hdr.cells[i]
        set_cell_width(cell, width)
        set_cell_margins(cell, top=90, bottom=90, start=75, end=75)
        shade_cell(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, bold=True, east_asia="黑体")
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": BORDER}, bottom={"val": "single", "sz": "8", "color": BORDER}, left={"val": "single", "sz": "8", "color": BORDER}, right={"val": "single", "sz": "8", "color": BORDER})
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, (value, width) in enumerate(zip(row_data, widths)):
            cell = row.cells[i]
            set_cell_width(cell, width)
            set_cell_margins(cell, top=80, bottom=80, start=75, end=75)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.alignment = (alignments[i] if alignments else (WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY))
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            set_cell_border(cell, top={"val": "single", "sz": "6", "color": BORDER}, bottom={"val": "single", "sz": "6", "color": BORDER}, left={"val": "single", "sz": "6", "color": BORDER}, right={"val": "single", "sz": "6", "color": BORDER})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover(doc):
    sec = doc.sections[0]
    configure_section(sec, cover=True)
    clear_header_footer(sec)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("项目编号：________________")
    set_run_font(r, size=12, east_asia="宋体")
    p.paragraph_format.space_after = Pt(55)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("中国机器人及人工智能创新大赛")
    set_run_font(r, size=22, bold=True, east_asia="黑体")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("查 新 报 告")
    set_run_font(r, size=30, bold=True, east_asia="黑体")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("（正式稿）")
    set_run_font(r, size=15, east_asia="宋体")
    p.paragraph_format.space_after = Pt(70)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table)
    labels = ["项目名称：", "项目作者：", "查新完成日期："]
    values = [PROJECT_NAME, "", REPORT_DATE]
    for row_idx, (label, value) in enumerate(zip(labels, values)):
        set_cell_width(table.cell(row_idx, 0), 3.25)
        set_cell_width(table.cell(row_idx, 1), 12.25)
        for c in table.rows[row_idx].cells:
            set_cell_margins(c, top=120, bottom=120, start=40, end=40)
            set_cell_border(c, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p1 = table.cell(row_idx, 0).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.first_line_indent = Pt(0)
        r = p1.add_run(label)
        set_run_font(r, size=12, east_asia="宋体")
        p2 = table.cell(row_idx, 1).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.first_line_indent = Pt(0)
        if value:
            r = p2.add_run(value)
            set_run_font(r, size=11.5, east_asia="宋体")
        else:
            r = p2.add_run("________________________________")
            set_run_font(r, size=11.5)


def add_toc(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("目  录")
    set_run_font(r, size=20, bold=True, east_asia="黑体")
    p.paragraph_format.space_after = Pt(18)
    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.first_line_indent = Pt(0)
    toc_p.paragraph_format.line_spacing = 1.35
    add_field(toc_p, 'TOC \\o "1-2" \\h \\z \\u', "目录将在打开文档时自动生成")
    doc.add_page_break()


def add_report_body(doc):
    add_section_heading(doc, "一、查新项目名称")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table)
    set_cell_width(table.cell(0, 0), 3.3)
    set_cell_width(table.cell(0, 1), 13.2)
    for idx, txt in enumerate(("查新项目名称", PROJECT_NAME)):
        cell = table.cell(0, idx)
        set_cell_margins(cell, top=170, bottom=170, start=100, end=100)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": BORDER}, bottom={"val": "single", "sz": "8", "color": BORDER}, left={"val": "single", "sz": "8", "color": BORDER}, right={"val": "single", "sz": "8", "color": BORDER})
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(txt)
        set_run_font(r, size=10.5, bold=(idx == 0), east_asia="黑体" if idx == 0 else "宋体")
    add_body_paragraph(doc, "项目定位：面向宁波港及一般海运出口场景，为中小出口企业、跨境卖家、货代与仓库装柜人员提供港前装柜决策、三维作业复核和合规预审支持。项目名称中的“数字孪生”指由装箱算法真实坐标驱动的三维作业过程映射，不表示与物理设备实时双向联动的工业级数字孪生。当前系统属于具备 SaaS 工作流形态、面向 SaaS 化交付的轻量化 Web 系统原型。", first_indent=True)

    add_section_heading(doc, "二、查新目的")
    add_body_paragraph(doc, "本报告用于报名参加第二十八届中国机器人及人工智能大赛人工智能创新赛。查新工作围绕项目的五个技术要点，检索国内外公开可访问的正式论文、国际规则、专利题录和商业系统资料，查证在所查范围内是否已有相同或相近研究，对检出资料与项目技术方案进行逐项比较，并对项目的新颖性作出审慎判断。")
    add_body_paragraph(doc, "查新结论服务于竞赛申报与技术说明，不替代法定科技查新机构意见，也不构成专利法意义上的新颖性、创造性、侵权或自由实施鉴定。")

    add_section_heading(doc, "三、查新项目的创新要点")
    add_body_paragraph(doc, "项目创新定位集中于面向出口装柜业务的工程化组合、系统集成与场景应用，不将已有的三维装箱、装船配载、数字孪生、规则检索或人工复核单项能力表述为底层理论首创。")
    innovation_rows = [
        ("1", "Fast Block多约束三维装箱引擎与六类业务约束协同机制", "以真实货物尺寸、重量、数量和业务属性为输入，采用确定性块优先策略，在 LIFO 卸货顺序、动态重心、同 SKU 聚集、堆叠安全、搬运可达和 VGM/总重合规六类业务约束下生成三维坐标方案。"),
        ("2", "SKU-柜型-装载结果一体化输出", "在同一后端链路中完成柜型比较，输出三维坐标、装载顺序、已装与未装清单及结果指标。Web 端现有导出范围为 HTML、CSV、JSON 与浏览器打印；小程序报告导出和后端原生 PDF 生成不列为完成能力。"),
        ("3", "真实坐标驱动的三维作业复核", "采用 18 相位状态机，将装箱坐标和顺序转化为可播放、暂停、跳步和复核的作业过程，表达门端、门洞、叉车、作业路径、碰撞与状态变化；该能力属于作业过程可视化，不作物理仿真或工业级数字孪生认定。"),
        ("4", "规则依据、AI 解释与人工复核相衔接", "依据 SOLAS/VGM、CTU Code、ISPM 15 等规则，在本地规则判断基础上采用 TF-IDF 法规检索和 AI 风险解释，并保留人工复核环节。AI 不替代装箱算法及本地规则结论；法规 PDF 原文、页码和高亮联动不列为完成能力。"),
        ("5", "面向中小出口主体的贯通式工作流", "形成“货物录入-SKU 管理-柜型比较-多约束装箱-三维复核-合规预审-AI 解释-人工复核-结果导出”的轻量化 Web 主链路。微信小程序承担部分方案查看能力，系统不作成熟多租户商业 SaaS 认定。"),
    ]
    add_table(doc, ["序号", "创新要点", "技术构成与事实边界"], innovation_rows, [1.15, 4.55, 10.8], font_size=9)
    add_subheading(doc, "演示指标说明", 2)
    demo_rows = [("NB2026-DEMO-01", "40HQ", "221 件", "221 件", "0 件", "84.51%", "本地规则 WARNING / AI建议 PASS / 人工复核独立")]
    add_table(doc, ["订单", "柜型", "输入", "装入", "未装", "空间利用率", "合规结果"], demo_rows, [3.1, 1.7, 1.7, 1.7, 1.7, 2.2, 2.4], font_size=8.5)
    add_body_paragraph(doc, "上述数据仅代表一次演示订单结果，不用于证明全部货物均已装入，也不外推为系统平均水平、稳定上限或普遍性能。", first_indent=True)

    add_section_heading(doc, "四、查新点")
    novelty_points = [
        "查新点一：面向装柜业务的 Fast Block多约束三维装箱引擎，采用六类业务约束协同机制，在确定性输出中同时处理 LIFO 卸货顺序、重心、SKU 聚集、堆叠安全、搬运可达和 VGM/总重合规。",
        "查新点二：基于真实 SKU 尺寸、重量、数量与业务属性进行柜型比较，并由同一后端生成三维坐标、装载顺序、未装货物清单和结果指标。",
        "查新点三：将装箱算法输出转化为可播放、暂停、跳步和复核的三维作业过程，以真实坐标和状态机表达门端、门洞、叉车、路径、碰撞与作业状态。",
        "查新点四：将 SOLAS/VGM、CTU Code、ISPM 15 等规则依据与本地规则判定、法规检索、AI 风险解释及人工复核串联为港前预审链路。",
        "查新点五：面向中小出口企业、跨境卖家、货代和仓库装柜人员，构建贯通 SKU 管理、柜型比较、三维装箱、三维复核、港前合规、人工复核和结果导出的轻量化工作流。",
    ]
    for item in novelty_points:
        add_bullet(doc, item)

    add_section_heading(doc, "五、文献检索范围及检索策略")
    add_subheading(doc, "（一）检索时点与时限", 2)
    add_body_paragraph(doc, "检索完成日期为 2026 年 7 月 11 日。检索时限以 2016 年至检索日为主，并补充 1993 年以来与三维装箱、集装箱配载、码头搬运、稳定性约束和装载规划有关的经典正式研究。")
    add_subheading(doc, "（二）检索来源", 2)
    source_rows = [
        ("正式论文", "出版方论文页、DOI 登记页", "装箱、配载、码头搬运、稳定性、承载与多港顺序"),
        ("国际规则", "IMO、ILO/UNECE、IPPC 官方页面及官方文件", "SOLAS/VGM、CTU Code、ISPM 15"),
        ("专利", "Google Patents 可访问题录页", "stowage plan、container transfer、container stacking"),
        ("商业系统", "Kaleris、Herbert-ABS、AnyLogic、Yseop 官方产品页", "装载计算、配载规划、三维表达、供应链仿真、受监管流程与人工控制"),
        ("补充文献", "arXiv 题录页", "多港配载、三维装箱、数字孪生、法律 RAG、人工在环；统一作为预印本列入附件"),
        ("项目事实", "项目事实核查文件、申报表述-证据矩阵及代码证据索引", "五个查新点、演示数据及未完成功能边界"),
    ]
    add_table(doc, ["资料类型", "检索来源", "主要主题"], source_rows, [2.5, 5.6, 8.4], font_size=9)
    add_subheading(doc, "（三）检索词", 2)
    terms_rows = [
        ("中文", "三维装箱；集装箱装载优化；块优先装箱；重心约束；支撑稳定；承载约束；多目的港；LIFO 卸货；装载顺序；集装箱配载；码头搬运；三维作业可视化；数字孪生港口；VGM；CTU Code；ISPM 15；法规检索；人工复核"),
        ("英文", "3D bin packing; container loading; cargo stability; load bearing; center of gravity; multi-drop; multi-port stowage; relocation; unloading sequence; loading computer; digital twin container terminal; verified gross mass; cargo transport unit; wood packaging; regulation retrieval; human-in-the-loop"),
    ]
    add_table(doc, ["语种", "检索词"], terms_rows, [2.0, 14.5], font_size=9)
    add_subheading(doc, "（四）代表性检索式", 2)
    strategies = [
        '("3D bin packing" OR "container loading") AND (stability OR support OR "load bearing" OR "center of gravity")',
        '("container loading" OR "bin packing") AND (multi-drop OR LIFO OR "unloading sequence")',
        '("container ship" AND stowage) AND (complexity OR relocation OR "master bay plan")',
        '("container terminal" AND handling) AND (storage OR relocation OR "digital twin")',
        '("loading computer" OR "stowage planning software") AND (stability OR VGM OR lashing)',
        '(SOLAS AND VGM) OR "verified gross mass of a packed container"',
        '"IMO/ILO/UNECE CTU Code" AND packing',
        '"ISPM 15" AND "wood packaging material"',
        '("regulation retrieval" OR "legal RAG") AND (audit OR compliance)',
        '("human-in-the-loop" OR "human oversight") AND compliance',
        '"Method for determining a stowage plan" AND patent',
        '(SKU AND "container type") AND (packing OR loading OR visualization)',
    ]
    for idx, strategy in enumerate(strategies, 1):
        add_body_paragraph(doc, f"{idx}. {strategy}", first_indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5)

    add_section_heading(doc, "六、检索结果")
    add_subheading(doc, "（一）相关资料检出情况", 2)
    add_body_paragraph(doc, "在所列来源与检索时限内，检出与本项目密切或较密切相关的正式论文 6 篇，国际规则依据 3 项，商业系统对照 4 项，专利题录 3 项；另检出多篇 arXiv 预印本和开源线索，统一列入补充附件。正文核心判断以正式论文、官方规则、商业产品官方页面及可核实专利题录为依据。")

    add_subheading(doc, "（二）经典与正式研究", 2)
    paper_rows = [
        ("P01", "Ambrosino, Sciomachen, Tanfani（2004），master bay plan", "以结构与作业限制建立船舶主贝位计划模型和启发式方法，说明集装箱配载、港序与作业限制已有成熟研究基础。", "2、5"),
        ("P02", "Taleb-Ibrahimi, de Castilho, Daganzo（1993），terminal storage vs handling", "研究出口集装箱在码头的堆存空间与搬运工作权衡，说明码头堆存、搬运与作业效率并非空白问题。", "3、5"),
        ("P03", "Avriel, Penn, Shpirer（2000），stowage complexity", "证明集装箱船配载中的倒箱/移箱问题具有组合复杂性，并与图着色问题相关。", "1、2"),
        ("P10", "Bischoff, Ratcliff（1995），practical container loading", "讨论稳定、均衡装载和多点卸货等实践因素，表明空间利用与业务约束的联合处理已有长期研究。", "1、2"),
        ("P11", "Bortfeldt, Wäscher（2013），constraints review", "系统归纳容器、货物、装载与实践约束，证明多约束集装箱装载已有广泛研究积累。", "1、2、5"),
        ("P12", "Junqueira, Morabito, Yamashita（2012），stability and load bearing", "建立含垂直/水平稳定与承载能力的三维装载模型，直接覆盖支撑、稳定、承载和易碎性等约束。", "1"),
    ]
    add_table(doc, ["编号", "资料", "主要内容及对比意义", "关联查新点"], paper_rows, [1.2, 5.15, 8.35, 1.8], font_size=8.5)

    add_subheading(doc, "（三）国际规则依据", 2)
    standard_rows = [
        ("STD01", "SOLAS 第 VI 章第 2 条及 VGM 配套文件", "IMO 规定托运人应提供经核验的载货集装箱总重，VGM 是装船条件之一，并应及时提供给船长/代表和码头代表用于配载计划。", "4"),
        ("STD02", "2014 IMO/ILO/UNECE CTU Code", "为海陆联运货物运输单元的装载、固定、危险货物和相关培训提供全球性实践准则，支撑装载安全与作业预审。", "1、4"),
        ("STD03", "ISPM 15：国际贸易中木质包装材料监管", "规定降低木质包装传播检疫性有害生物风险的处理与标识要求，支撑木质包装合规检查。", "4"),
    ]
    add_table(doc, ["编号", "规则依据", "与项目的关系", "关联查新点"], standard_rows, [1.3, 5.1, 8.3, 1.8], font_size=9)

    add_subheading(doc, "（四）商业和系统对照", 2)
    product_rows = [
        ("CP01", "Kaleris MACS3 / StowMan", "MACS3 覆盖船舶装载计算、稳性、强度、系固及规则集成；StowMan 面向配载规划和装载能力提升。", "说明船舶级装载计算与配载规划已有成熟产品；项目差异在港前 SKU 装柜、三维作业复核和合规预审的轻量化组合。"),
        ("CP02", "Herbert-ABS CargoMax", "支持装载规划、稳性与强度评估、集装箱详细录入、BAPLIE、系固、IMDG 分隔检查和三维图形。", "说明装载安全、限制检查和三维表达已有工业产品；项目不以这些单点能力作为首创。"),
        ("CP03", "AnyLogic / anyLogistix", "支持供应链与物流仿真、优化、动态场景分析和数字孪生建模。", "说明物流数字孪生与过程仿真已有成熟平台；项目差异在装箱坐标驱动的现场作业复核场景。"),
        ("CP04", "Yseop", "面向受监管环境提供结构化内容生成、可追溯性、审计记录、质量检查和人工控制。", "说明受监管流程中的 AI 生成、追溯与人工审核已有相邻产品；项目差异在出口装柜规则链路。"),
    ]
    add_table(doc, ["编号", "产品/系统", "公开能力", "对比意义"], product_rows, [1.2, 3.5, 5.4, 6.4], font_size=8.5)

    add_subheading(doc, "（五）专利题录", 2)
    add_body_paragraph(doc, "PAT01：Method for determining a stowage plan，US5809489A（仅核实题名及专利号）。该题录表明自动生成配载计划、考虑重量分布和倒箱成本的技术思路已有专利公开；鉴于正文仅按题名及专利号保守使用，不以该条目单独承担项目新颖性结论。", first_indent=True)
    add_body_paragraph(doc, "PAT02 与 PAT03 仅作为历史性集装箱搬运/堆叠线索列入附件，不作为核心新颖性判断依据。", first_indent=True)

    add_subheading(doc, "（六）五个查新点的逐项分析与结论", 2)

    analyses = [
        (
            "1. 查新点一：Fast Block多约束三维装箱引擎与六类业务约束协同机制",
            "检出情况：P10、P11、P12 已分别讨论稳定、均衡、多点卸货、实践约束、支撑与承载；P03 说明配载与移箱问题具有长期研究基础；CP02 也覆盖装载限制、稳性、系固与三维表达。",
            "项目特征：项目将 LIFO 卸货顺序、动态重心、同 SKU 聚集、堆叠安全、搬运可达和 VGM/总重合规六类约束嵌入面向装柜业务的确定性求解链路，并将结果继续用于三维作业复核与合规预审。",
            "比较结论：已有资料能够否定“三维装箱底层理论空白”或“单项约束首次提出”的表述；在所查范围内，项目的可辨识差异主要体现在六类约束面向具体出口装柜业务的协同实现，以及算法结果与后续复核、预审链路的连接。该点判断为工程化改进与组合创新。",
        ),
        (
            "2. 查新点二：真实 SKU、柜型比较与一体化结果输出",
            "检出情况：P01 对船舶主贝位计划的结构和作业限制进行建模；PAT01 涉及自动确定配载计划；CP01、CP02 已覆盖船舶装载数据、配载规划、稳性/强度和装载限制等能力。",
            "项目特征：项目面向中小出口主体，以 SKU 尺寸、重量、数量和业务属性为输入，在轻量化 Web 流程中比较柜型，由同一后端形成三维坐标、装载顺序、已装/未装清单与结果指标。NB2026-DEMO-01 仅为一次 40HQ 演示订单：221 件输入、221 件装入、0 件未装、空间利用率 84.51%、总重 8,812 kg；本地规则 WARNING、AI建议 PASS、人工复核保持独立。",
            "比较结论：柜位规划、装载结果和未装处理均有先例，项目差异不在首次输出三维结果，而在 SKU 录入、柜型比较、坐标、顺序、未装清单与轻量化结果导出的贯通。现有输出边界为 Web 端 HTML、CSV、JSON 与浏览器打印；小程序导出和后端原生 PDF 生成不列为完成能力。该点判断为工程化改进与场景应用创新。",
        ),
        (
            "3. 查新点三：真实坐标驱动的三维作业过程复核",
            "检出情况：P02 已研究码头堆存与搬运工作；CP03 表明供应链、物流仿真与数字孪生建模已有成熟平台；CP02 也具备三维图形表达。",
            "项目特征：项目使用算法输出的真实三维坐标和装载顺序驱动 18 相位状态机，形成门端、门洞、叉车、作业路径、碰撞和状态变化的可播放、暂停、跳步、复核过程。其目标是支持装柜作业理解与复核，而非对物理设备进行实时数据回传、双向控制或高保真动力学仿真。",
            "比较结论：数字孪生、三维可视化和物流仿真并非新概念；项目差异主要表现为将装箱坐标、作业角色、路径与状态机组合为面向出口装柜人员的过程复核工具。该点判断为场景应用创新与系统集成创新。",
        ),
        (
            "4. 查新点四：国际规则、AI 风险解释与人工复核的港前预审",
            "检出情况：STD01、STD02、STD03 分别为 VGM、运输单元装载实践和木质包装提供明确规则根基；CP01、CP02 已将部分海事规则和限制检查纳入装载系统；CP04 表明受监管环境中的可追溯生成、审计和人工控制已有工业实践。",
            "项目特征：项目采用本地规则结论、TF-IDF 法规检索、AI 风险解释和人工复核相衔接的链路，AI 仅承担建议与解释，不替代装箱算法或本地规则判断。当前法规证据以结构化文本为主，法规 PDF 原文、页码和命中区域高亮联动不列为完成能力。",
            "比较结论：VGM、CTU Code、ISPM 15 及人工在环合规均有既有基础；项目的可辨识差异在于把这些能力嵌入出口装柜港前预审并与装箱结果联动。该点不宜认定为合规 AI 单项首创，判断为组合创新。",
        ),
        (
            "5. 查新点五：面向中小出口主体的贯通式轻量化工作流",
            "检出情况：P01-P03、P10-P12 分别覆盖配载、搬运、复杂性和实践约束；CP01、CP02 强在船舶级装载计算与配载，CP03 强在仿真和数字孪生，CP04 强在受监管内容与人工控制。现有公开资料多覆盖完整链路中的一段或若干段。",
            "项目特征：项目以中小出口企业、跨境卖家、货代和仓库装柜人员为对象，形成从 SKU 管理、柜型比较、三维装箱到三维作业复核、港前合规、AI 解释、人工复核和 Web 端结果导出的主链路。当前为轻量化 Web 系统原型，微信小程序部分功能受限，不作成熟多租户 SaaS 认定。",
            "比较结论：截至检索日，在可访问公开资料中，暂未检出能够明确证明同一系统完整覆盖上述全部链路的资料；但该判断受到中文数据库访问范围、商业产品非公开功能和专利全文可见度限制。该点判断为系统集成创新与场景应用创新。",
        ),
    ]
    for title, found, project, conclusion in analyses:
        add_subheading(doc, title, 3)
        add_body_paragraph(doc, found, bold_lead="检出情况：")
        add_body_paragraph(doc, project, bold_lead="项目特征：")
        add_body_paragraph(doc, conclusion, bold_lead="比较结论：")

    add_section_heading(doc, "七、查新结论", page_break_before=True)
    add_body_paragraph(doc, "经对检出资料进行分析和比较，可以确认：三维装箱、集装箱配载、稳定性、支撑、承载、重心、堆叠、多点卸货、港序与倒箱等问题均已有正式研究；船舶装载计算、配载规划、三维图形、供应链数字孪生和受监管流程中的可追溯 AI 也已有商业系统；SOLAS/VGM、CTU Code 和 ISPM 15 为相关合规检查提供了既有国际规则依据。")
    add_body_paragraph(doc, "因此，本项目不属于三维装箱底层理论或上述单项能力的首创。项目较为稳妥的创新定位是：在宁波港及一般海运出口场景中，围绕中小出口企业、跨境卖家、货代与仓库装柜人员，形成“多约束装箱-三维作业复核-港前合规预审-AI 风险解释-人工复核-结果导出”的工程化组合与系统集成。")
    add_body_paragraph(doc, "对五个查新点的结论分别为：查新点一为工程化改进与组合创新；查新点二为工程化改进与场景应用创新；查新点三为场景应用创新与系统集成创新；查新点四为组合创新；查新点五为系统集成创新与场景应用创新。")
    add_body_paragraph(doc, "截至 2026 年 7 月 11 日，在本报告所列公开可访问范围内，暂未检出能够明确证明同一公开系统完整覆盖上述五点及其贯通工作流的资料。该结论具有检索范围和证据可见度边界，不作绝对排他性判断。")


def add_statement_and_attachments(doc):
    doc.add_page_break()
    add_section_heading(doc, "八、申报者本人、所在学院签字盖章的查新声明与证明")
    add_body_paragraph(doc, "1. 本报告中陈述的项目事实是真实和准确的。", first_indent=False)
    add_body_paragraph(doc, "2. 我们按照大赛查新规范进行检索、文献分析和审核，并在所查范围内作出上述查新结论。", first_indent=False)

    sig = doc.add_table(rows=2, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig.autofit = False
    set_table_fixed(sig)
    for i in range(2):
        set_cell_width(sig.cell(0, i), 8.25)
        set_cell_margins(sig.cell(0, i), top=180, bottom=120, start=180, end=180)
        set_cell_border(sig.cell(0, i), top={"val": "single", "sz": "8", "color": BORDER}, bottom={"val": "single", "sz": "8", "color": BORDER}, left={"val": "single", "sz": "8", "color": BORDER}, right={"val": "single", "sz": "8", "color": BORDER})
        sig.cell(0, i).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    sig.rows[0].height = Cm(7.8)
    sig.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    labels = ["申报者（签字）：", "申报者所在学院（盖章）："]
    for i, label in enumerate(labels):
        p = sig.cell(0, i).paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(label)
        set_run_font(r, size=11, bold=True, east_asia="黑体")
    date_cell = sig.cell(1, 0).merge(sig.cell(1, 1))
    set_cell_margins(date_cell, top=160, bottom=160, start=180, end=180)
    set_cell_border(date_cell, top={"val": "single", "sz": "8", "color": BORDER}, bottom={"val": "single", "sz": "8", "color": BORDER}, left={"val": "single", "sz": "8", "color": BORDER}, right={"val": "single", "sz": "8", "color": BORDER})
    p = date_cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("日期：________年____月____日")
    set_run_font(r, size=11, east_asia="宋体")

    doc.add_page_break()
    add_section_heading(doc, "九、附件清单")
    attachment_rows = [
        ("附件一", "核心正式研究、规则、商业系统与 PAT01 题录", "列明 P01、P02、P03、P10、P11、P12、STD01-STD03、CP01-CP04、PAT01 的正式题录及官方链接"),
        ("附件二", "补充文献与开源线索", "arXiv 资料统一标注为预印本，不作为正式期刊成果；用于补充相邻技术发展"),
        ("附件三", "补充专利题录", "PAT02、PAT03 仅作历史性线索，不作为核心新颖性判断依据"),
        ("附件四", "项目事实与能力边界", "列明演示订单适用范围、系统形态和未完成能力，供申报审核"),
    ]
    add_table(doc, ["编号", "附件名称", "内容说明"], attachment_rows, [2.0, 5.4, 9.1], font_size=9)

    add_section_heading(doc, "十、备注（检索限制与适用范围）")
    limitations = [
        "检索受到 CNKI、万方、维普及部分专利全文站点访问条件限制；中文数据库中的未公开或不可访问记录可能未被覆盖。",
        "商业系统具有非公开功能、版本差异和客户定制属性，官方公开页面不能穷尽其全部能力。",
        "百科、搜索结果片段及其他二级公开页面仅用于线索定位，不作为正文正式引用来源；正式题录和核心功能描述以出版方、标准制定机构、专利题录页或产品官方页面为准。",
        "arXiv 资料均以“预印本”列入附件，不按正式期刊或正式会议成果使用，也不单独承担核心新颖性判断。",
        "本报告中的“暂未检出相同公开系统”仅限于所列检索范围和检索时点，不等同于专利法意义上的新颖性、创造性、侵权或自由实施鉴定。",
    ]
    for idx, text in enumerate(limitations, 1):
        add_body_paragraph(doc, f"{idx}. {text}", first_indent=False)


def add_reference_entry(doc, code, citation, summary, url, link_label="访问来源"):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.hanging_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"[{code}] {citation}")
    set_run_font(r, size=9.5)
    if url:
        r = p.add_run("  ")
        set_run_font(r, size=9.5)
        add_hyperlink(p, link_label, url)
    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Pt(18)
    p2.paragraph_format.line_spacing = 1.2
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run(f"摘要说明：{summary}")
    set_run_font(r, size=9)


def add_appendices(doc):
    add_section_heading(doc, "附件一  核心正式研究、规则、商业系统与 PAT01 题录", page_break_before=True)
    add_subheading(doc, "A. 正式研究", 2)
    refs = [
        ("P01", "Ambrosino D, Sciomachen A, Tanfani E. Stowing a containership: the master bay plan problem[J]. Transportation Research Part A: Policy and Practice, 2004, 38(2): 81-99. DOI: 10.1016/j.tra.2003.09.002.", "研究主贝位计划问题，以结构和作业限制建立模型与启发式方案，并以实际码头实例验证。", "https://doi.org/10.1016/j.tra.2003.09.002"),
        ("P02", "Taleb-Ibrahimi M, de Castilho B, Daganzo C F. Storage space vs handling work in container terminals[J]. Transportation Research Part B: Methodological, 1993, 27(1): 13-32. DOI: 10.1016/0191-2615(93)90009-Y.", "研究出口集装箱码头堆存策略、空间需求与搬运工作量之间的权衡。", "https://doi.org/10.1016/0191-2615(93)90009-Y"),
        ("P03", "Avriel M, Penn M, Shpirer N. Container ship stowage problem: complexity and connection to the coloring of circle graphs[J]. Discrete Applied Mathematics, 2000, 103(1-3): 271-279. DOI: 10.1016/S0166-218X(99)00245-0.", "研究集装箱船配载中倒箱成本最小化问题，论证其组合复杂性。", "https://doi.org/10.1016/S0166-218X(99)00245-0"),
        ("P10", "Bischoff E E, Ratcliff M S W. Issues in the development of approaches to container loading[J]. Omega, 1995, 23(4): 377-390. DOI: 10.1016/0305-0483(95)00015-G.", "讨论稳定、均衡装载、多点卸货与实践约束，说明基本空间模型之外还需处理多类业务因素。", "https://doi.org/10.1016/0305-0483(95)00015-G"),
        ("P11", "Bortfeldt A, Wäscher G. Constraints in container loading - A state-of-the-art review[J]. European Journal of Operational Research, 2013, 229(1): 1-20. DOI: 10.1016/j.ejor.2012.12.006.", "系统回顾集装箱装载中的容器、货物与装载约束，指出实践约束已有丰富研究。", "https://doi.org/10.1016/j.ejor.2012.12.006"),
        ("P12", "Junqueira L, Morabito R, Yamashita D S. Three-dimensional container loading models with cargo stability and load bearing constraints[J]. Computers & Operations Research, 2012, 39(1): 74-85. DOI: 10.1016/j.cor.2010.07.017.", "建立含垂直/水平稳定和承载能力的三维装载模型，覆盖支撑、承载及易碎性等实践因素。", "https://doi.org/10.1016/j.cor.2010.07.017"),
    ]
    for item in refs:
        add_reference_entry(doc, *item, link_label="DOI/出版方页面")

    add_subheading(doc, "B. 国际规则", 2)
    standards = [
        ("STD01", "International Maritime Organization. Verification of the gross mass of a packed container[EB/OL]. SOLAS regulation VI/2; resolution MSC.380(94); MSC.1/Circ.1475.", "明确托运人提供 VGM 的责任、VGM 对配载计划的用途及其作为装船条件的地位。", "https://www.imo.org/en/ourwork/safety/pages/verification-of-the-gross-mass.aspx"),
        ("STD02", "International Maritime Organization, International Labour Organization, United Nations Economic Commission for Europe. IMO/ILO/UNECE Code of Practice for Packing of Cargo Transport Units (CTU Code)[S]. 2014. MSC.1/Circ.1497.", "为货物运输单元的装载、固定、危险货物和相关作业提供全球实践准则。", "https://www.imo.org/en/ourwork/safety/pages/ctu-code.aspx"),
        ("STD03", "International Plant Protection Convention. ISPM 15: Regulation of wood packaging material in international trade[S/OL].", "规定降低木质包装材料传播检疫性有害生物风险的处理和标识要求。", "https://www.ippc.int/en/publications/640/"),
    ]
    for item in standards:
        add_reference_entry(doc, *item, link_label="官方页面")

    add_subheading(doc, "C. 商业和系统对照", 2)
    products = [
        ("CP01", "Kaleris. MACS3 Suite / StowMan[EB/OL].", "MACS3 提供船舶装载计算、稳性、强度、系固与规则集成；StowMan 面向配载规划。", "https://kaleris.com/ncvssafety/"),
        ("CP02", "Herbert-ABS. CargoMax for Container Ships[EB/OL].", "公开能力包括装载规划、稳性/强度、详细集装箱录入、BAPLIE、系固、IMDG 分隔检查和三维图形。", "https://www.herbert-abs.com/cargomax-for-container-ships"),
        ("CP03", "The AnyLogic Company. Supply Chain Simulation Software / anyLogistix[EB/OL].", "面向供应链和物流仿真、优化、动态场景分析与数字孪生建模。", "https://www.anylogic.com/supply-chains/"),
        ("CP04", "Yseop. Regulatory-grade AI for submission-ready documents[EB/OL].", "公开强调结构化生成、可追溯性、审计记录、质量检查和人工控制，属于受监管流程的相邻系统。", "https://yseop.com/why-yseop/"),
    ]
    for item in products:
        add_reference_entry(doc, *item, link_label="产品官方页面")

    add_subheading(doc, "D. PAT01", 2)
    add_reference_entry(doc, "PAT01", "Method for determining a stowage plan. US5809489A[P].", "仅核实题名及专利号。正文仅据此确认自动确定配载计划的早期专利线索，不依赖未完整核验字段作强结论。", "https://patents.google.com/patent/US5809489A/en", link_label="专利题录页")

    add_section_heading(doc, "附件二  补充文献与开源线索（预印本）", page_break_before=True)
    add_body_paragraph(doc, "本附件所列 arXiv 资料统一作为预印本使用，未按正式期刊或正式会议成果列入正文核心证据。其作用是说明相邻技术方向的近期发展，不单独承担项目新颖性结论。")
    preprints = [
        ("P04", "Catarina Junqueira 等（2020）", "An Integrated Optimization Model for the Multi-Port Stowage Planning and the Container Relocation Problems", "arXiv:2006.06795；预印本", "1、2"),
        ("P05", "Faouzi Hakimi 等（2025）", "Towards a Digital Twin Modeling Method for Container Terminal Port", "arXiv:2503.13511；预印本", "3、5"),
        ("P06", "Neel Guha 等（2023）", "LegalBench: A Collaboratively Built Benchmark for Measuring Legal Reasoning in Large Language Models", "arXiv:2308.11462；预印本来源", "4"),
        ("P07", "Nicholas Pipitone, Ghita Houir Alami（2024）", "LegalBench-RAG: A Benchmark for Retrieval-Augmented Generation in the Legal Domain", "arXiv:2408.10343；预印本", "4"),
        ("P08", "Haitao Li 等（2025）", "LexRAG: Benchmarking Retrieval-Augmented Generation in Multi-Turn Legal Consultation Conversation", "arXiv:2502.20640；预印本", "4"),
        ("P09", "Maurice Chiodo 等（2025）", "Formalising Human-in-the-Loop: Computational Reductions, Failure Modes, and Legal-Moral Responsibility", "arXiv:2505.10426；预印本", "4"),
        ("CN06", "Hang Zhao 等（2020）", "Online 3D Bin Packing with Constrained Deep Reinforcement Learning", "arXiv:2006.14978；预印本", "1"),
        ("CN07", "Qiruyi Zuo 等（2022）", "A Constructive Heuristic Algorithm for 3D Bin Packing of Irregular Shaped Items", "arXiv:2206.15116；预印本", "1"),
        ("CN08", "Lei Gao 等（2025）", "One4Many-StablePacker: An Efficient Deep Reinforcement Learning Framework for the 3D Bin Packing Problem", "arXiv:2510.10057；预印本", "1"),
        ("CN09", "Xijun Li 等（2021）", "Learning to Optimize Industry-Scale Dynamic Pickup and Delivery Problems", "arXiv:2105.12899；预印本来源", "5"),
        ("CN10", "Naiyao Wang 等（2024）", "KUNPENG: An Embodied Large Model for Intelligent Maritime", "arXiv:2407.09048；预印本", "3、5"),
        ("OS03", "Heng Xiong 等（2024）", "GOPT: Generalizable Online 3D Bin Packing via Transformer-based Deep Reinforcement Learning", "arXiv:2409.05344；预印本/代码线索", "1、3"),
    ]
    table = add_table(doc, ["编号", "作者/年份", "题名", "来源与状态", "关联点"], preprints, [1.2, 3.3, 6.5, 3.8, 1.7], font_size=7.8)
    arxiv_urls = {
        "P04": "https://arxiv.org/abs/2006.06795", "P05": "https://arxiv.org/abs/2503.13511", "P06": "https://arxiv.org/abs/2308.11462", "P07": "https://arxiv.org/abs/2408.10343",
        "P08": "https://arxiv.org/abs/2502.20640", "P09": "https://arxiv.org/abs/2505.10426", "CN06": "https://arxiv.org/abs/2006.14978", "CN07": "https://arxiv.org/abs/2206.15116",
        "CN08": "https://arxiv.org/abs/2510.10057", "CN09": "https://arxiv.org/abs/2105.12899", "CN10": "https://arxiv.org/abs/2407.09048", "OS03": "https://arxiv.org/abs/2409.05344",
    }
    add_body_paragraph(doc, "预印本链接：", first_indent=False, bold_lead="预印本链接：", size=9.5)
    for code, url in arxiv_urls.items():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"[{code}] ")
        set_run_font(r, size=9)
        add_hyperlink(p, url, url)

    add_section_heading(doc, "附件三  补充专利题录", page_break_before=True)
    patent_rows = [
        ("PAT02", "Shipboard freight container transferring apparatus", "US3042227A", "附件线索；不作为核心新颖性判断依据", "https://patents.google.com/patent/US3042227A/en"),
        ("PAT03", "Container stacking system", "US3231103A", "附件线索；不作为核心新颖性判断依据", "https://patents.google.com/patent/US3231103A/en"),
    ]
    table = add_table(doc, ["编号", "题名", "专利号", "使用方式", "来源"], [(a, b, c, d, "专利题录页") for a, b, c, d, _ in patent_rows], [1.4, 5.0, 2.4, 5.2, 2.5], font_size=8.5)
    for row_idx, (_, _, _, _, url) in enumerate(patent_rows, 1):
        p = table.cell(row_idx, 4).paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_hyperlink(p, "官方题录", url)
    add_body_paragraph(doc, "PAT02 与 PAT03 仅反映历史性集装箱搬运和堆叠装置线索，与项目五个查新点的算法-三维复核-合规预审主线对应较弱，故不进入正文核心新颖性判断。")

    add_section_heading(doc, "附件四  项目事实与能力边界")
    boundaries = [
        ("算法表述", "Fast Block多约束三维装箱引擎；六类业务约束协同机制", "用于描述 LIFO、重心、SKU 聚集、堆叠安全、搬运可达和 VGM/总重合规的协同处理。"),
        ("演示订单", "NB2026-DEMO-01：221 件输入、221 件装入、0 件未装、84.51%、8,812 kg", "仅代表一次演示订单；不外推为系统平均水平、稳定上限或普遍性能。"),
        ("系统形态", "具备 SaaS 工作流形态、面向 SaaS 化交付的轻量化 Web 系统原型", "不作成熟多租户商业 SaaS 认定。"),
        ("三维能力", "真实坐标驱动的三维装柜作业过程可视化", "不作工业级数字孪生或物理仿真认定；不包含物理设备实时数据回传与双向控制。"),
        ("法规证据", "结构化法规文本、本地规则、TF-IDF 检索、AI 风险解释", "法规 PDF 原文、页码与命中区域高亮联动不列为完成能力。"),
        ("导出能力", "Web 端 HTML、CSV、JSON 与浏览器打印", "微信小程序报告导出和后端原生 PDF 生成不列为完成能力。"),
        ("AI 角色", "风险解释、法规检索与辅助建议", "AI 不替代确定性装箱算法及本地规则结论。"),
        ("人工复核", "保留人工复核环节和记录", "不扩展表述为成熟企业级审批流或完整多租户治理体系。"),
    ]
    add_table(doc, ["事项", "可写事实", "适用边界"], boundaries, [3.0, 6.5, 7.0], font_size=8.8)


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "甬柜易算项目科技查新报告（正式稿）"
    props.subject = "第二十八届中国机器人及人工智能大赛人工智能创新赛查新报告"
    props.author = ""
    props.keywords = "甬柜易算; 科技查新; 三维装箱; 港前合规; 人工智能创新赛"
    props.comments = "A4 正式查新报告；签字、盖章和日期均留空。"


def add_document_settings(doc):
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def validate_document(doc):
    texts = []
    for p in doc.paragraphs:
        texts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    all_text = "\n".join(texts)
    forbidden = ["本次修订", "此前底稿", "证据清洗", "六维加权" + "算法", "cite", "filecite", "Wikipedia", "维基百科", "‐", "‑", "‒", "–", "—"]
    found = [x for x in forbidden if x in all_text]
    if found:
        raise ValueError(f"发现禁用或过程性文本：{found}")
    required = [
        "Fast Block多约束三维装箱引擎", "六类业务约束协同机制", "NB2026-DEMO-01", "221 件装入", "0 件未装", "84.51%",
        "仅代表一次演示订单", "成熟多租户", "法规 PDF", "小程序报告导出", "后端原生 PDF", "仅核实题名及专利号",
        "P01", "P02", "P03", "STD01", "STD02", "STD03", "CP01", "CP02", "CP03", "CP04", "预印本",
    ]
    missing = [x for x in required if x not in all_text]
    if missing:
        raise ValueError(f"缺少必需事实或证据标签：{missing}")


def build():
    doc = Document()
    configure_styles(doc)
    set_document_properties(doc)
    add_document_settings(doc)
    add_cover(doc)

    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_section, cover=False)
    add_running_header(main_section)
    add_running_footer(main_section)
    set_page_number_start(main_section, 1)

    add_toc(doc)
    add_report_body(doc)
    add_statement_and_attachments(doc)
    add_appendices(doc)

    validate_document(doc)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
