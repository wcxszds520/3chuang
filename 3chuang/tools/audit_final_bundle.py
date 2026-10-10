from __future__ import annotations

import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageChops
try:
    from pypdf import PdfReader
except ModuleNotFoundError:
    from PyPDF2 import PdfReader


ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "output" / "final-submission"
QA = ROOT / "output" / "project-report" / "qa-final"


def docx_full_text(path: Path) -> tuple[str, Document]:
    doc = Document(path)
    chunks: list[str] = []
    chunks.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        chunks.extend(p.text for p in section.header.paragraphs)
        chunks.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(chunks), doc


def section_sizes(doc: Document):
    return [
        {
            "width_cm": round(section.page_width.cm, 2),
            "height_cm": round(section.page_height.cm, 2),
            "orientation": "landscape" if section.page_width > section.page_height else "portrait",
        }
        for section in doc.sections
    ]


def pdf_summary(path: Path, render_folder: Path | None = None):
    reader = PdfReader(path)
    sizes = []
    for page in reader.pages:
        w = float(page.mediabox.width)
        h = float(page.mediabox.height)
        sizes.append({"w": round(w, 1), "h": round(h, 1), "orientation": "landscape" if w > h else "portrait"})
    render = {"count": 0, "blank_pages": [], "edge_touch_pages": []}
    if render_folder and render_folder.exists():
        images = sorted(render_folder.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
        render["count"] = len(images)
        for index, file in enumerate(images, 1):
            with Image.open(file) as image:
                gray = image.convert("L")
                diff = ImageChops.difference(gray, Image.new("L", gray.size, 255))
                bbox = diff.point(lambda value: 255 if value > 14 else 0).getbbox()
                if bbox is None:
                    render["blank_pages"].append(index)
                    continue
                left, top, right, bottom = bbox
                if left <= 3 or top <= 3 or right >= gray.width - 3 or bottom >= gray.height - 3:
                    render["edge_touch_pages"].append(index)
    return {"pages": len(reader.pages), "page_sizes": sizes, "render": render}


def seq_field_counts(docx_path: Path):
    with zipfile.ZipFile(docx_path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    values = [node.text or "" for node in root.findall(".//w:instrText", ns)]
    types = [values[i + 1] for i, value in enumerate(values[:-1]) if value.strip() == "SEQ"]
    return {"figure": types.count("图"), "table": types.count("表"), "total": len(types)}


def occurrence_context(text: str, term: str):
    out = []
    for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE):
        start = max(0, match.start() - 45)
        end = min(len(text), match.end() + 45)
        out.append(re.sub(r"\s+", " ", text[start:end]))
    return out


def main():
    main_docx = FINAL / "01_甬柜易算_项目研究报告.docx"
    prompt_docx = FINAL / "04_甬柜易算_研究报告图像生成提示词.docx"
    screenshot_docx = FINAL / "06_甬柜易算_真实系统截图补录清单.docx"
    novelty_docx = FINAL / "08_甬柜易算_科技查新报告.docx"
    main_text, main_doc = docx_full_text(main_docx)
    prompt_text, prompt_doc = docx_full_text(prompt_docx)
    screenshot_text, screenshot_doc = docx_full_text(screenshot_docx)
    novelty_text, novelty_doc = docx_full_text(novelty_docx)

    chapter_headings = [p.text for p in main_doc.paragraphs if p.style.name == "Heading 1" and re.match(r"^第[一二三四五六七八九十]+章", p.text)]
    table_titles = [p.text for p in main_doc.paragraphs if p.style.name == "Table Caption"]
    fig_titles = [p.text for p in main_doc.paragraphs if p.style.name == "Figure Caption"]
    all_text = "\n".join([main_text, prompt_text, screenshot_text, novelty_text])
    forbidden = [
        "六维加权" + "算法",
        "44" + "4/" + "86" + "4",
        "60" + ".6%",
        "WARNING" + "/86",
        "本次修订",
        "此前底稿",
        "证据清洗",
    ]
    scans = {term: occurrence_context(all_text, term) for term in forbidden}

    novelty_terms = {
        term: occurrence_context(novelty_text, term)
        for term in ["本次修订", "此前底稿", "证据清洗", "PAT01", "PAT02", "PAT03", "arXiv", "预印本", "Wikipedia"]
    }

    workbook = load_workbook(FINAL / "05_甬柜易算_研究报告引用与证据核对表.xlsx", read_only=True)
    with zipfile.ZipFile(FINAL / "03_甬柜易算_研究报告图像素材包.zip") as archive:
        zip_names = archive.namelist()
    pngs = [name for name in zip_names if name.lower().endswith(".png")]
    generated_pngs = [name for name in pngs if name.startswith("01_GPT_Image_2/")]
    real_pngs = [name for name in pngs if name.startswith("02_真实系统截图/")]

    result = {
        "main_docx": {
            "sections": section_sizes(main_doc),
            "tables": len(main_doc.tables),
            "table_captions": len(table_titles),
            "seq_fields": seq_field_counts(main_docx),
            "figure_captions_total": len(fig_titles),
            "inline_shapes": len(main_doc.inline_shapes),
            "chapters": len(chapter_headings),
            "chapter_titles": chapter_headings,
            "references_1_to_13_present": all(f"[{i}]" in main_text for i in range(1, 14)),
            "required_demo_values_present": all(
                value in main_text
                for value in ["NB2026-DEMO-01", "221 件", "0 件", "84.51%", "8,812 kg"]
            ),
        },
        "prompt_docx": {
            "sections": section_sizes(prompt_doc),
            "tables": len(prompt_doc.tables),
            "inline_shapes": len(prompt_doc.inline_shapes),
            "spec_headings": len([p for p in prompt_doc.paragraphs if p.style.name == "Heading 2" and re.match(r"^图\d+ ", p.text)]),
        },
        "screenshot_docx": {
            "sections": section_sizes(screenshot_doc),
            "tables": len(screenshot_doc.tables),
            "has_13_rows": all(f"S{i:02d}" in screenshot_text for i in range(1, 14)),
        },
        "novelty_docx": {
            "sections": section_sizes(novelty_doc),
            "tables": len(novelty_doc.tables),
            "terms": novelty_terms,
            "has_five_points": all(f"查新点{n}" in novelty_text for n in "一二三四五"),
            "blank_signature_lines": "申报者签字：________________________________" in novelty_text and "日期：______年____月____日" in novelty_text,
        },
        "pdfs": {
            "main": pdf_summary(FINAL / "01_甬柜易算_项目研究报告.pdf", QA / "main"),
            "evidence": pdf_summary(FINAL / "02_甬柜易算_项目研究佐证材料汇编.pdf", QA / "evidence"),
            "novelty": pdf_summary(FINAL / "08_甬柜易算_科技查新报告.pdf", QA / "novelty"),
            "quality": pdf_summary(FINAL / "07_甬柜易算_研究报告质量检查记录.pdf", QA / "quality"),
        },
        "workbook": {"sheets": workbook.sheetnames},
        "zip": {
            "entries": len(zip_names),
            "pngs": len(pngs),
            "generated_pngs": len(generated_pngs),
            "real_pngs": len(real_pngs),
            "has_legacy": any("legacy" in name.lower() for name in zip_names),
            "has_prompt_doc": any("提示词" in name for name in zip_names),
            "has_specs": any("规格表" in name for name in zip_names),
            "has_sources": any("来源说明" in name for name in zip_names),
        },
        "forbidden_scan": scans,
    }
    QA.mkdir(parents=True, exist_ok=True)
    out = QA / "audit.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
